"""
build_study_guide_docx.py
-------------------------
Render docs/STUDY_GUIDE.md as a .docx in the same house style as the bilingual
companion, but for a monolingual document that carries equations.

It differs from build_explainer_docx.py in what it has to handle. The study
guide is a derivation, so fenced blocks hold aligned equations rather than
shell commands: they are set in a monospace box with the line breaks kept,
because an equation reflowed to the paragraph width is unreadable. It also
handles bullet and numbered lists, and blockquotes, which the companion's
source never used.

Usage:
    python scripts/build_study_guide_docx.py

Output:
    docs/STUDY_GUIDE.docx
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent.parent
SRC = ROOT / "docs" / "STUDY_GUIDE.md"
OUT = ROOT / "docs" / "STUDY_GUIDE.docx"

BLUE = RGBColor(0x0F, 0x4C, 0x81)
NAVY = RGBColor(0x16, 0x21, 0x3E)
GREY = RGBColor(0x44, 0x44, 0x44)
BODY = RGBColor(0x33, 0x33, 0x33)
CODE_TXT = RGBColor(0x1F, 0x2D, 0x3D)

FILL_CODE = "F4F6F8"    # equations and commands
FILL_KEY = "E8F1FA"     # blockquote call-outs
NAVY_HEX = "16213E"


def shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def strip_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        borders.append(e)
    table._tbl.tblPr.append(borders)


def inline(par, text, *, size=10.5, color=BODY, bold=False):
    """Write text into a paragraph, honouring **bold**, *italic* and `code`."""
    for part in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)", text):
        if not part:
            continue
        run = par.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text, run.bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(size - 0.5)
            run.font.color.rgb = color
            continue
        elif part.startswith("*") and part.endswith("*"):
            run.text, run.italic = part[1:-1], True
        else:
            run.text = part
            run.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"


def add_box(doc, lines, fill, *, mono=False, size=10.5):
    """A tinted single-cell box. Monospace boxes keep their line breaks."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    strip_borders(t)
    cell = t.cell(0, 0)
    shade(cell, fill)
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0 if mono else 3)
        if mono:
            # An equation must not reflow, so it is written as one literal run
            # with its spacing intact rather than passed through inline().
            r = p.add_run(line if line.strip() else " ")
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = CODE_TXT
        else:
            inline(p, line, size=size, color=BODY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, NAVY_HEX)
        c.paragraphs[0].text = ""
        inline(c.paragraphs[0], h, size=9.5,
               color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    for r in body:
        cells = t.add_row().cells
        for i, v in enumerate(r[:len(header)]):
            cells[i].paragraphs[0].text = ""
            inline(cells[i].paragraphs[0], v, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_list_item(doc, text, *, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.8)
    inline(p, text)


def main():
    md = SRC.read_text().split("\n")
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = BODY
    st.paragraph_format.space_after = Pt(6)

    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.1)
        s.top_margin = s.bottom_margin = Cm(1.9)
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("dMRI Rosetta Stone — Study Guide")
        r.font.size = Pt(8)
        r.font.color.rgb = GREY

    i = 0
    first_heading = True
    while i < len(md):
        line = md[i].rstrip()

        if not line:
            i += 1
            continue

        # Fenced block: equations or commands, line breaks preserved.
        if line.lstrip().startswith("```"):
            i += 1
            block = []
            while i < len(md) and not md[i].lstrip().startswith("```"):
                block.append(md[i].rstrip())
                i += 1
            i += 1
            while block and not block[0].strip():
                block.pop(0)
            while block and not block[-1].strip():
                block.pop()
            if block:
                add_box(doc, block, FILL_CODE, mono=True)
            continue

        if line.startswith("|"):
            rows, i = parse_table(md, i)
            if rows:
                add_table(doc, rows)
            continue

        if line.startswith("---"):
            i += 1
            continue

        # Blockquote: the "★" call-outs. Collected as one tinted box.
        if line.startswith(">"):
            block, para = [], []
            while i < len(md) and md[i].lstrip().startswith(">"):
                stripped = md[i].lstrip()[1:].strip()
                if stripped:
                    para.append(stripped)
                else:
                    if para:
                        block.append(" ".join(para))
                        para = []
                i += 1
            if para:
                block.append(" ".join(para))
            if block:
                add_box(doc, block, FILL_KEY)
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(23)
            r.font.color.rgb = BLUE
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if line.startswith("## "):
            # No page break before the first section, which follows the title.
            if not first_heading:
                doc.add_page_break()
            first_heading = False
            p = doc.add_paragraph()
            r = p.add_run(line[3:])
            r.bold = True
            r.font.size = Pt(16)
            r.font.color.rgb = NAVY
            p.paragraph_format.space_after = Pt(8)
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph()
            r = p.add_run(line[4:])
            r.bold = True
            r.font.size = Pt(12.5)
            r.font.color.rgb = BLUE
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if re.match(r"^[-*] ", line):
            add_list_item(doc, line[2:].strip())
            i += 1
            continue

        if re.match(r"^\d+\. ", line):
            add_list_item(doc, re.sub(r"^\d+\.\s*", "", line), numbered=True)
            i += 1
            continue

        # Ordinary paragraph: join its lines, stopping at any block opener.
        block = []
        while i < len(md) and md[i].strip() and not md[i].lstrip().startswith(
                ("#", "|", "---", ">", "```", "- ", "* ")) and not \
                re.match(r"^\d+\. ", md[i].lstrip()):
            block.append(md[i].strip())
            i += 1
        if block:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            inline(p, " ".join(block))

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
